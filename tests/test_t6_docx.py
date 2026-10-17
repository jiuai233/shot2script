"""T6 的生成、读回和实际 OOXML 破坏测试。"""

from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree as ET
import pytest

from s2s.docx_build import TEMPLATE, W, build_docx, read_parts, serialize, text_of, xml
from s2s.docx_check import check_docx

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests/fixtures/t6/三集合并稿.md"
TITLE = "T6样例 剧本合集"


@pytest.fixture
def generated(tmp_path):
    path = tmp_path / "样例.docx"
    build_docx(FIXTURE, path, TITLE)
    return path


def rules(path, total=3, missing=None):
    return {error["规则"] for error in check_docx(path, total, missing)["错误"]}


def rewrite(path, change):
    parts = read_parts(path)
    document = xml(parts["word/document.xml"])
    change(document)
    parts["word/document.xml"] = serialize(document)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def find_paragraph(doc, starts):
    return next(p for p in doc.findall(W + "body/" + W + "p") if text_of(p).startswith(starts))


def test_three_episode_fixture_passes(generated):
    result = check_docx(generated, 3)
    assert result["通过"], result
    assert result["收录集号"] == [1, 2, 3]
    assert result["一级大纲数"] == 3
    assert result["二级大纲数"] == result["场景数"] == 6


def test_round_trip_is_line_exact(generated):
    texts = [p.text for p in Document(generated).paragraphs]
    assert texts[:2] == [TITLE, ""]
    assert texts[2:] == FIXTURE.read_text(encoding="utf-8").splitlines()


def test_all_episode_blanks_use_template_spacing(generated):
    doc = xml(read_parts(generated)["word/document.xml"])
    for n in range(1, 4):
        blank = find_paragraph(doc, f"第{n}集").getprevious()
        assert not text_of(blank)
        spacing = blank.find(W + "pPr/" + W + "spacing")
        assert spacing.get(W + "before") == spacing.get(W + "after") == "240"


def test_reuses_template_parts_and_section(generated):
    original, built = read_parts(TEMPLATE), read_parts(generated)
    assert original.keys() == built.keys()
    for name in original:
        if name != "word/document.xml":
            assert original[name] == built[name], name
    old_section = xml(original["word/document.xml"]).find(W + "body/" + W + "sectPr")
    section = xml(built["word/document.xml"]).find(W + "body/" + W + "sectPr")
    assert ET.tostring(section, method="c14n") == ET.tostring(old_section, method="c14n")
    assert all(p.getparent().tag == W + "body" for p in xml(built["word/document.xml"]).iter(W + "p"))


@pytest.mark.parametrize("old,new,rule", [
    ("第2集", "第4集", "集号"),
    ("场景1-1：会议室，日，内", "场景1-1 会议室 日 内", "场景格式"),
    ("场景1-2：走廊", "场景1-3：走廊", "场景序号"),
    ("场景2-1：会议室", "场景1-1：会议室", "场景所属集"),
    ("人物：林舟、沈宁、穿红衣服的小男孩", "△ 林舟打开文件。", "人物行"),
])
def test_bad_markdown_reports_specific_issue(tmp_path, old, new, rule):
    source = tmp_path / "错误稿.md"
    source.write_text(FIXTURE.read_text(encoding="utf-8").replace(old, new, 1), encoding="utf-8")
    out = tmp_path / "错误.docx"
    build_docx(source, out, TITLE)
    assert rule in rules(out)


def test_missing_episode_semantics(tmp_path):
    lines = FIXTURE.read_text(encoding="utf-8")
    # 全剧共三集，第二集尚未收录；参数仍传总集数 3。
    source = tmp_path / "缺集.md"
    source.write_text(lines[:lines.index("第2集")] + lines[lines.index("第3集"):], encoding="utf-8")
    out = tmp_path / "缺集.docx"
    build_docx(source, out, TITLE)
    assert "集号" in rules(out)
    result = check_docx(out, 3, {2})
    assert result["通过"], result
    assert result["一级大纲数"] == 2
    assert "缺失集出现" in rules(out, 3, {1})


@pytest.mark.parametrize("numbers", [[1, 1, 3], [1, 3, 2], [1, 2, 4]])
def test_duplicate_unordered_or_extra_episodes(generated, numbers):
    def change(doc):
        ps = [find_paragraph(doc, f"第{n}集") for n in range(1, 4)]
        for p, n in zip(ps, numbers):
            p.find(W + "r/" + W + "t").text = f"第{n}集"
    rewrite(generated, change)
    assert "集号" in rules(generated)


@pytest.mark.parametrize("starts,tag,value,rule", [
    ("第1集", "outlineLvl", "2", "导航层级"),
    ("场景1-1", "outlineLvl", "4", "导航层级"),
    ("人物：", "outlineLvl", "0", "一级大纲数量"),
    ("人物：", "outlineLvl", "1", "二级大纲数量"),
    ("人物：", "pStyle", "2", "段落样式"),
    (TITLE, "jc", "left", "大标题"),
])
def test_actual_paragraph_properties(generated, starts, tag, value, rule):
    def change(doc):
        prop = find_paragraph(doc, starts).find(W + "pPr/" + W + tag)
        prop.set(W + "val", value)
    rewrite(generated, change)
    assert rule in rules(generated)


@pytest.mark.parametrize("starts,tag,slot,value,rule", [
    ("人物：", "rFonts", "eastAsia", "黑体", "字体"),
    ("场景1-1", "rFonts", "ascii", "Calibri", "字体"),
    ("第1集", "rFonts", "hAnsi", "Arial", "字体"),
    (TITLE, "rFonts", "cs", "Arial", "字体"),
    ("人物：", "rFonts", "asciiTheme", "minorHAnsi", "字体"),
    ("人物：", "sz", "val", "28", "字号"),
    ("场景1-1", "szCs", "val", "24", "字号"),
    (TITLE, "b", "val", "0", "加粗"),
])
def test_actual_run_properties(generated, starts, tag, slot, value, rule):
    def change(doc):
        node = find_paragraph(doc, starts).find(W + "r/" + W + "rPr/" + W + tag)
        node.set(W + slot, value)
    rewrite(generated, change)
    assert rule in rules(generated)


def test_missing_run_font_or_size_does_not_use_theme(generated):
    def change(doc):
        rpr = find_paragraph(doc, "人物：").find(W + "r/" + W + "rPr")
        for tag in ("rFonts", "sz", "szCs"):
            rpr.remove(rpr.find(W + tag))
    rewrite(generated, change)
    assert {"字体", "字号"} <= rules(generated)


def test_effective_outline_resolves_style_inheritance(generated):
    def change(doc):
        ppr = find_paragraph(doc, "第1集").find(W + "pPr")
        ppr.remove(ppr.find(W + "outlineLvl"))
    rewrite(generated, change)
    # 范本 heading 3 的继承值为 2，不能冒充一级导航。
    assert {"导航层级", "一级大纲数量"} <= rules(generated)


def test_episode_requires_blank_before_it(generated):
    def change(doc):
        p = find_paragraph(doc, "第1集")
        p.getparent().remove(p.getprevious())
    rewrite(generated, change)
    assert "集前空段" in rules(generated)


def test_scene_person_line_must_be_immediate(generated):
    def change(doc):
        p = find_paragraph(doc, "场景1-1")
        p.addnext(ET.Element(W + "p"))
    rewrite(generated, change)
    assert "人物行" in rules(generated)


@pytest.mark.parametrize("damage", ["outside_body", "missing_section", "nested_body", "page_size", "missing_body"])
def test_document_structure(generated, damage):
    def change(doc):
        body = doc.find(W + "body")
        if damage == "outside_body":
            doc.append(deepcopy(find_paragraph(doc, "人物：")))
        elif damage == "missing_section":
            body.remove(body.find(W + "sectPr"))
        elif damage == "nested_body":
            ET.SubElement(body, W + "body")
        elif damage == "page_size":
            body.find(W + "sectPr/" + W + "pgSz").set(W + "w", "1")
        else:
            doc.remove(body)
    rewrite(generated, change)
    assert ("页面设置" if damage == "page_size" else "文档结构") in rules(generated)


def test_corrupt_package_reports_json_error(tmp_path, generated):
    invalid = tmp_path / "损坏.docx"
    invalid.write_bytes(b"not a document")
    assert "文档结构" in rules(invalid)
    assert "文档结构" in rules(tmp_path / "不存在.docx")
    parts = read_parts(generated)
    with ZipFile(generated, "w") as archive:
        for name, data in parts.items():
            if name != "word/styles.xml":
                archive.writestr(name, data)
    assert "文档结构" in rules(generated)


def test_text_is_not_rewritten_and_tabs_survive(tmp_path):
    source = tmp_path / "原文.md"
    lines = ["第1集", "场景1-1：会议室，日，内", "人物：林舟", "林舟：  <合同> & 〔附件〕\t原件。  "]
    source.write_text("\n".join(lines) + "\n", encoding="utf-8")
    out = tmp_path / "原文.docx"
    build_docx(source, out, TITLE)
    assert [p.text for p in Document(out).paragraphs][2:] == lines
    assert check_docx(out, 1)["通过"]


def test_builder_rejects_empty_source_and_overwriting_input(tmp_path):
    source = tmp_path / "空.md"
    source.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="不能为空"):
        build_docx(source, tmp_path / "空.docx", TITLE)
    with pytest.raises(ValueError, match="不能覆盖"):
        build_docx(source, source, TITLE)


@pytest.mark.parametrize("count,missing", [(0, set()), (3, {0}), (3, {4})])
def test_bad_parameters(generated, count, missing):
    assert "参数" in rules(generated, count, missing)


def test_cli_success_and_failure_json(tmp_path):
    output = tmp_path / "命令行.docx"
    build = subprocess.run([sys.executable, str(ROOT / "tools/docx_build.py"), str(FIXTURE), str(output), TITLE],
                           cwd=tmp_path, capture_output=True, text=True)
    assert build.returncode == 0, build.stderr
    assert json.loads(build.stdout)["通过"]
    for count, flag in [("3", True), ("4", False)]:
        check = subprocess.run([sys.executable, str(ROOT / "tools/docx_check.py"), str(output), count],
                               cwd=tmp_path, capture_output=True, text=True)
        assert check.returncode == (0 if flag else 1), check.stderr
        assert json.loads(check.stdout)["通过"] is flag
    for csv in ["1", "错误"]:
        check = subprocess.run([sys.executable, str(ROOT / "tools/docx_check.py"), str(output), "3", "--缺失集", csv],
                               capture_output=True, text=True)
        assert check.returncode == 1, check.stderr
        assert not json.loads(check.stdout)["通过"]
